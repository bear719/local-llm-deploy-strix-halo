#!/usr/bin/env python3
"""渐进式披露（Progressive Disclosure）文档代理 + OpenAI 兼容透传。

两种模式（环境变量 PROXY_MODE 控制，默认 passthrough）：
  passthrough  —— 所有 /v1/* 请求原样透传到上游（含 stream SSE、/v1/models），
                 不注入 system prompt、不加强制工具。适合 WorkBuddy / OpenWebUI 等
                 自带上下文管理的客户端接入（等价于直连上游的兼容代理）。
  disclosure   —— 非流式 POST /v1/chat/completions 走"渐进式披露"工具循环
                 （注入目录 + 按需 search/peek）；流式请求仍透传（流式下无法做循环）。

客户端把大文件注册进来后，在 disclosure 模式下模型每轮只看"目录"（元数据 + 少量预览），
需要细节时由模型自己调用工具按需读取：
    list_files()                  列出已注册文件
    peek_file(id, start, stop)    按行区间读取（字符级近似分块）
    search_file(id, query)        关键词定位，返回命中行 ±2 行
    extract_file(id)              结构摘要（标题树 / 表格轮廓 / 页目录）
    load_file(id, max_chars)      全文兜底（默认截断 8000 字符）

用法:
    python3 docproxy.py serve  [--port 8081] [--upstream http://127.0.0.1:8080]
    PROXY_MODE=disclosure python3 docproxy.py serve   # 开启渐进式披露
    python3 docproxy.py add <文件路径>...        # 注册文件（写入索引）
    python3 docproxy.py ls                       # 查看已注册文件
    python3 docproxy.py ask "问题"               # 快速自测一轮披露对话（非流式）

预加载策略（disclosure 模式写入 system prompt 的初始可见量）:
    纯文本/md:  前 2 KB
    PDF/DOCX:   仅元数据 + 第一页/前 500 字
    图片:       仅元数据
"""
import argparse
import http.client
import json
import os
import re
import select
import socket
import sys
import time
import urllib.parse
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

PY = sys.executable
UPSTREAM = os.environ.get("UPSTREAM", "http://127.0.0.1:8080")
PROXY_MODE = os.environ.get("PROXY_MODE", "passthrough").lower()

# 默认给转发请求注入 chat_template_kwargs.enable_thinking=false。
# 背景（2026-10-02 实测）：Qwen3.6-35B-A3B 的模板默认开思考链，通用 OpenAI 客户端
# （如 OpenClaw，baseUrl 指向本代理）不会传 enable_thinking，结果 64 个 token 全被
# reasoning_content 吃掉、content 为空 —— 客户端拿到空回复。
# 注意模板差异：Qwen3.6 系列只认 enable_thinking，不认 /no_think 后缀；
# 而 qwen4exp（Qwen3.8-Flash-Next）正好相反。所以统一用 chat_template_kwargs 这个字段。
# 客户端若显式传了 enable_thinking，以客户端为准，不覆盖。
INJECT_NO_THINK = os.environ.get("INJECT_NO_THINK", "1").lower() not in ("0", "false", "no")

# 上游单次请求总超时（秒）。2026-10-02 修复：原值 1800 太长，
# 上游卡死时客户端要干等 30 分钟才报错，改为 600 尽快失败让客户端重试。
UPSTREAM_TIMEOUT = int(os.environ.get("UPSTREAM_TIMEOUT", "600"))
# 流式转发时「单次 read 的 socket 超时」（秒）。必须 >> 最长 prefill 耗时：
# 本机实测 84K tokens 冷启动 prefill ≈ 150s，更长上下文可达数百秒。
# ⚠ 断不可直接套用下面的 HEARTBEAT_SEC！2026-10-03 事故根因：
#   CPython 的 socket 一旦触发一次 socket.timeout，会把内部 _timed_out 永久置位，
#   此后任何 read 都立刻抛 OSError("cannot read from timed out object")（它是普通
#   OSError，不被 except (socket.timeout, TimeoutError) 捕获），连接直接报废 →
#   返回 502 → 客户端期待 SSE 却收到明文错误 → "aborted"。
#   表现特征：耗时稳定卡在超时值、上游状态码=200、首字节=None、字节=0。
# → 心跳改用 select 轮询实现（不置位 _timed_out）， socket 超时只作整体兜底。
STREAM_READ_TIMEOUT = int(os.environ.get("STREAM_READ_TIMEOUT", "600"))
# 流式转发时上游无数据的心跳间隔（秒）：定期发 SSE 注释，
# 避免长 prefill 期间被客户端/中间层的 idle 超时判定为死连接。
# 该值只用于 select 轮询间隔，严禁赋给 socket.settimeout（见上方说明）。
HEARTBEAT_SEC = int(os.environ.get("HEARTBEAT_SEC", "15"))
# 访问日志开关：默认开，排障时可看首字节/总耗时/上游状态码。
LOG_REQUESTS = os.environ.get("LOG_REQUESTS", "1").lower() not in ("0", "false", "no")
INDEX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docindex.json")
PREVIEW_LIMITS = {".md": 2000, ".txt": 2000, ".log": 2000, ".csv": 1500,
                  ".pdf": 500, ".docx": 500, ".pptx": 500, ".xlsx": 800}


# ---------- 文本抽取（只读，不改动源文件） ----------

def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".txt", ".log", ".csv", ".json", ".py", ".sh", ".html"):
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext == ".pdf":
        import pymupdf
        doc = pymupdf.open(path)
        return "\n\n".join(f"[第 {i+1} 页]\n" + pg.get_text()
                           for i, pg in enumerate(doc))
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for tb in d.tables:
            for row in tb.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[工作表 {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    return None  # 图片等二进制：只给元数据


def structure_outline(text, limit=1200):
    """标题树 / 结构轮廓，代替全文进入目录。"""
    lines = text.splitlines()
    out = []
    for i, ln in enumerate(lines, 1):
        if re.match(r"^#{1,4} ", ln) or re.match(r"^\[\d+页?\]|^\[第 \d+ 页\]", ln):
            out.append(f"  L{i}: {ln.strip()[:80]}")
        if len(out) >= 40:
            out.append("  ...(标题过多已截断)")
            break
    if not out:  # 无标题结构 → 给段落首行采样
        for i, ln in enumerate(lines, 1):
            if ln.strip() and i % max(1, len(lines) // 20) == 0:
                out.append(f"  L{i}: {ln.strip()[:60]}")
    return "\n".join(out[:40]) or "  (无可见结构)"


# ---------- 文件索引 ----------

def load_index():
    if os.path.exists(INDEX):
        with open(INDEX, encoding="utf-8") as f:
            return json.load(f)
    return {"next_id": 1, "files": {}}


def save_index(idx):
    with open(INDEX, "w", encoding="utf-8") as f:
        json.dump(idx, f, ensure_ascii=False, indent=1)


def register(path):
    path = os.path.abspath(path)
    if not os.path.isfile(path):
        print(f"[ERROR] 文件不存在: {path}")
        return None
    idx = load_index()
    st = os.stat(path)
    ext = os.path.splitext(path)[1].lower()
    meta = {"path": path, "name": os.path.basename(path), "ext": ext,
            "size_kb": round(st.st_size / 1024, 1)}
    text = extract_text(path)
    if text is None:
        meta.update(kind="binary", chars=0, lines=0)
        meta["preview"] = "(二进制文件，仅提供元数据)"
    else:
        meta.update(kind="text", chars=len(text), lines=text.count("\n") + 1,
                    preview=text[:PREVIEW_LIMITS.get(ext, 2000)])
    fid = str(idx["next_id"]); idx["next_id"] += 1
    idx["files"][fid] = meta
    save_index(idx)
    print(f"[OK] id={fid} {meta['name']}  {meta['size_kb']} KB  "
          f"{meta.get('chars', 0)} 字符")
    return fid


def get_text_cached(fid, _cache={}):
    if fid not in _cache:
        meta = load_index()["files"].get(fid)
        if not meta:
            return None
        _cache[fid] = extract_text(meta["path"])
    return _cache[fid]


# ---------- 工具实现（模型调用） ----------

def tool_call(name, args):
    idx = load_index()["files"]
    if name == "list_files":
        rows = [f"{k}: {v['name']} ({v['size_kb']} KB, {v.get('chars',0)} 字符, "
                f"{v['lines']} 行)" for k, v in idx.items()]
        return "\n".join(rows) or "(没有已注册文件)"

    fid = str(args.get("id", ""))
    meta = idx.get(fid)
    if not meta:
        return f"[ERROR] 文件 id {fid} 不存在，先调用 list_files"

    if name == "peek_file":
        text = get_text_cached(fid) or ""
        lines = text.splitlines()
        start = max(1, int(args.get("start", 1)))
        stop = min(len(lines), int(args.get("stop", start + 100)))
        if start > len(lines):
            return f"[ERROR] start 超出范围（共 {len(lines)} 行）"
        chunk = "\n".join(lines[start-1:stop])
        return chunk[:8000] + ("\n...(本段截断)" if len(chunk) > 8000 else "")

    if name == "search_file":
        text = get_text_cached(fid) or ""
        q = str(args.get("query", "")).strip()
        if not q:
            return "[ERROR] query 为空"
        lines = text.splitlines()
        hits, n = [], 0
        for i, ln in enumerate(lines):
            if q.lower() in ln.lower():
                ctx = "\n".join(lines[max(0, i-2):i+3])
                hits.append(f"--- L{i+1} ---\n{ctx}")
                n += 1
                if n >= 8:
                    hits.append("...(命中过多，仅前 8 处)")
                    break
        return "\n".join(hits) if hits else f"未找到 “{q}”"

    if name == "extract_file":
        text = get_text_cached(fid)
        if text is None:
            return f"(二进制文件) 名称={meta['name']} 大小={meta['size_kb']} KB"
        head = structure_outline(text)
        return (f"文件: {meta['name']}  共 {meta['lines']} 行 / {meta['chars']} 字符\n"
                f"结构轮廓:\n{head}")

    if name == "load_file":
        text = get_text_cached(fid) or ""
        cap = int(args.get("max_chars", 8000))
        return text[:cap] + ("\n...(截断，可用 peek_file 续读)" if len(text) > cap else "")

    return f"[ERROR] 未知工具 {name}"


TOOLS_SPEC = [{
    "type": "function", "function": {
        "name": n, "description": d,
        "parameters": {"type": "object", "properties": p, "required": r},
    },
} for n, d, p, r in [
    ("list_files", "列出所有已注册文件的 id、名称、大小、行数", {}, []),
    ("peek_file", "按行区间读取文件内容（一次最多约 8000 字符）",
     {"id": {"type": "string", "description": "文件 id"},
      "start": {"type": "integer", "description": "起始行，从 1 开始"},
      "stop": {"type": "integer", "description": "结束行（含）"}},
     ["id", "start", "stop"]),
    ("search_file", "在文件内搜索关键词，返回命中行及上下文（最多 8 处）",
     {"id": {"type": "string"}, "query": {"type": "string"}}, ["id", "query"]),
    ("extract_file", "返回文件的结构摘要：标题树 / 页目录 / 表格轮廓",
     {"id": {"type": "string"}}, ["id"]),
    ("load_file", "加载文件全文（默认截断 8000 字符，慎用，优先 peek/search）",
     {"id": {"type": "string"}, "max_chars": {"type": "integer"}}, ["id"]),
]]


def build_catalog():
    files = load_index()["files"]
    if not files:
        return ""
    parts = ["# 已注册文件（渐进式披露：以下只有目录，正文需调用工具获取）"]
    for fid, m in files.items():
        parts.append(f"\n## id={fid}  {m['name']}  ({m['size_kb']} KB, "
                     f"{m.get('chars', 0)} 字符, {m['lines']} 行)")
        parts.append(f"结构轮廓:\n{structure_outline(m['preview'])}" if m["kind"] == "text"
                     else "(二进制，仅元数据)")
        if m.get("preview"):
            p = m["preview"][:PREVIEW_LIMITS.get(m["ext"], 2000)]
            parts.append(f"开头预览:\n{p}")
    return "\n".join(parts)


SYSTEM_PROMPT = """你在按需提供文档内容的工作模式下工作。
你**没有**文档全文——只有下面的目录（元数据 + 结构轮廓 + 少量开头预览）。
回答涉及文档内容的问题时：
1. 先用 search_file 定位相关行，再用 peek_file 精读该区间；不要盲目 load_file。
2. 一次 peek 不超过 150 行；多个独立检索可以在同一轮并行发起。
3. 目录信息足够回答时（如"文档标题是什么"），直接回答，不调工具。

{catalog}"""


# ---------- media marker 净化（修复 mmproj tokenizer 400） ----------

# 2026-10-03 修复：llama-server 的 mmproj tokenizer 在请求文本里出现
# <__media_xxx__> 占位符但请求未附带任何图片（bitmaps=0）时，会直接 400
# "Failed to tokenize prompt"（日志 mtmd_tokenize: number of media markers
# in text (N) exceeds number of bitmaps (0)）。根因是客户端（如 WorkBuddy
# 上下文注入）在纯文本请求里残留了历史图片轮次留下的占位符但未附图。
# 本函数做防御性净化：当整条请求都不含任何图片时，剥离这些孤儿占位符，
# 让请求能正常走纯文本 tokenize；若请求确实带了图片，则保持不动交上游。
MEDIA_MARKER_RE = re.compile(r"<__media_[A-Za-z0-9_]+__>")


def _request_has_images(body):
    """整条请求是否携带任何图片内容（image_url / images / bitmaps / 内联 data:image）。"""
    if not isinstance(body, dict):
        return False
    if body.get("images") or body.get("bitmaps"):
        return True
    for m in body.get("messages", []) or []:
        c = m.get("content")
        if isinstance(c, list):
            for part in c:
                if isinstance(part, dict) and part.get("type") in ("image_url", "image"):
                    url = (part.get("image_url") or {}).get("url")
                    if url:
                        return True
        elif isinstance(c, str):
            if c.lstrip().startswith("data:image"):
                return True
    return False


def strip_stray_media_markers(body):
    """剥离孤儿 media 占位符。返回是否改动过 body。"""
    if not isinstance(body, dict):
        return False
    if _request_has_images(body):
        return False  # 真多模态请求，不动，交上游正常处理
    msgs = body.get("messages")
    if not isinstance(msgs, list):
        return False
    changed = False
    for m in msgs:
        c = m.get("content")
        if isinstance(c, str):
            new = MEDIA_MARKER_RE.sub("", c)
            if new != c:
                m["content"] = new
                changed = True
        elif isinstance(c, list):
            for part in c:
                if isinstance(part, dict) and part.get("type") == "text":
                    t = part.get("text", "")
                    if isinstance(t, str):
                        new = MEDIA_MARKER_RE.sub("", t)
                        if new != t:
                            part["text"] = new
                            changed = True
    if changed and LOG_REQUESTS:
        sys.stderr.write("[docproxy] 已剥离孤儿 media 占位符（纯文本请求误带视觉 marker）\n")
    return changed


# ---------- OpenAI 代理 ----------

def upstream_chat(payload):
    req = urllib.request.Request(
        f"{UPSTREAM}/v1/chat/completions",
        data=json.dumps(payload).encode(),
        headers={"Content-Type": "application/json"})
    return json.loads(urllib.request.urlopen(req, timeout=1200).read())


def handle_chat(body, verbose=True):
    """渐进式披露：非流式工具循环。"""
    messages = body.pop("messages", [])
    sysmsg = {"role": "system",
              "content": SYSTEM_PROMPT.format(catalog=build_catalog())}
    msgs = [sysmsg] + messages
    stats = {"tool_calls": [], "prompt_tokens": 0, "rounds": 0}
    for _ in range(8):  # tool 循环上限
        payload = dict(body, messages=msgs, tools=TOOLS_SPEC, tool_choice="auto")
        resp = upstream_chat(payload)
        choice = resp["choices"][0]
        stats["prompt_tokens"] += resp["usage"]["prompt_tokens"]
        stats["rounds"] += 1
        if choice["finish_reason"] != "tool_calls":
            resp["_stats"] = stats
            return resp
        msgs.append(choice["message"])
        for tc in choice["message"]["tool_calls"]:
            fn = tc["function"]["name"]
            try:
                args = json.loads(tc["function"]["arguments"] or "{}")
            except json.JSONDecodeError:
                args = {}
            result = tool_call(fn, args)
            stats["tool_calls"].append(f"{fn}({args})")
            if verbose:
                print(f"  [tool] {fn}({args}) -> {len(result)} 字符")
            msgs.append({"role": "tool", "tool_call_id": tc["id"],
                         "content": result})
    resp["_stats"] = stats
    return resp


# ---------- 透传（兼容 OpenAI 客户端） ----------

def forward_to_upstream(self, method, path, raw=None, stream=False):
    """把请求原样转发到上游并回传响应。stream=True 时逐块透传 SSE。

    2026-10-02 修复（流式响应永不结束，导致客户端一直"思考中"）：
      1. 旧实现手动发 `Connection: keep-alive`，会让 BaseHTTPRequestHandler
         把 close_connection 置为 False；而 protocol_version 仍是 HTTP/1.0、
         又无 Content-Length / chunked，客户端既收不到长度也等不到 EOF，
         只能挂到超时。现改为 HTTP/1.1 + Transfer-Encoding: chunked，
         并显式写 0 长度终止块，响应边界明确。
      2. 旧实现流式下恒返回 200，上游 503/500 的错误 JSON 被当成 SSE 转出，
         客户端解析不到合法 data 行同样悬挂。现按上游状态码分支。
      3. 上游长时间无数据（长 prefill）时发 SSE 注释心跳，防 idle 超时。
    """
    u = urllib.parse.urlparse(UPSTREAM)
    conn = http.client.HTTPConnection(u.hostname, u.port, timeout=UPSTREAM_TIMEOUT)
    body = raw if raw is not None else b""
    t0 = time.time()
    ttfb = None
    status = None
    headers_sent = False
    nbytes = 0
    try:
        conn.request(method, path, body=body,
                     headers={"Content-Type": "application/json"})
        resp = conn.getresponse()
        ctype = resp.getheader("Content-Type", "application/json")
        status = resp.status
        if stream and 200 <= status < 300:
            self.send_response(200)
            self.send_header("Content-Type", ctype or "text/event-stream")
            self.send_header("Cache-Control", "no-cache")
            self.send_header("X-Accel-Buffering", "no")
            self.send_header("Transfer-Encoding", "chunked")
            self.end_headers()
            headers_sent = True
            # socket 超时只作整体兜底（长 prefill 容忍），心跳靠下面的 select 轮询，
            # 不得把 HEARTBEAT_SEC 设给 settimeout——详见常量区注释。
            conn.sock.settimeout(STREAM_READ_TIMEOUT)
            deadline = time.time() + STREAM_READ_TIMEOUT
            fd = conn.sock.fileno()
            while True:
                ready, _, _ = select.select([fd], [], [], HEARTBEAT_SEC)
                if ready:
                    chunk = resp.read(4096)
                    if not chunk:
                        break
                    if ttfb is None:
                        ttfb = time.time() - t0
                    nbytes += len(chunk)
                    if not self._write_chunk(chunk):
                        break
                else:
                    # 上游暂无数据（长 prefill 中）：发 SSE 注释心跳保活
                    if time.time() > deadline:
                        break
                    if not self._write_chunk(b": keep-alive\n\n"):
                        break
            self._write_chunk(b"", last=True)
        else:
            data = resp.read()
            ttfb = time.time() - t0
            nbytes = len(data)
            # 上游非 2xx 时原样回状态码和错误体，不再伪装成 200
            self.send_response(status)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
    except Exception as e:
        if LOG_REQUESTS:
            sys.stderr.write(f"[docproxy] 转发异常 {method} {path}: {e!r} "
                             f"耗时={time.time()-t0:.1f}s\n")
        # 若流式响应头已发出（已处于 chunked 流中），不能再写 502 明文响应，
        # 否则半截 HTTP 头会污染 SSE 流，客户端 chunked 解析崩溃 → "aborted"。
        # 此时只能补 chunked 终止块并关闭连接。
        if headers_sent:
            try:
                self._write_chunk(b"", last=True)
            except Exception:
                pass
        else:
            try:
                self._send(502, {"error": f"upstream unreachable: {e}"})
            except Exception:
                pass
    finally:
        conn.close()
        if LOG_REQUESTS:
            sys.stderr.write(
                f"[docproxy] {method} {path} stream={int(bool(stream))} "
                f"上游={status} 首字节={ttfb if ttfb is None else round(ttfb,2)}s "
                f"总耗时={time.time()-t0:.2f}s 字节={nbytes}\n")


class Handler(BaseHTTPRequestHandler):
    # 关键：HTTP/1.1 才能正确使用 chunked / keep-alive 语义。
    # 保持 HTTP/1.0 时，流式响应若无长度字段，客户端只能靠 EOF 判定结束。
    protocol_version = "HTTP/1.1"

    def log_message(self, fmt, *a):
        if LOG_REQUESTS:
            sys.stderr.write("[%s] %s\n" % (time.strftime("%H:%M:%S"), fmt % a))

    def _write_chunk(self, data, last=False):
        """写一块 chunked 数据；客户端已断开返回 False。last=True 写终止块。"""
        try:
            if last:
                self.wfile.write(b"0\r\n\r\n")
            else:
                self.wfile.write(b"%x\r\n" % len(data) + data + b"\r\n")
            self.wfile.flush()
            return True
        except (BrokenPipeError, ConnectionResetError, OSError):
            return False

    def _send(self, code, obj):
        data = json.dumps(obj, ensure_ascii=False).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _read_body(self):
        """读取请求体，兼容 Content-Length 与 Transfer-Encoding: chunked 两种上传。

        背景（2026-10-03 修复）：WorkBuddy 等客户端在请求体较大（40K+ tokens≈100KB+）
        时会改用 chunked 分块上传且不带 Content-Length。原实现只认 Content-Length，
        chunked 时被读成空体 → 转发空 JSON → 上游 500 → 客户端（期待 SSE）报 aborted。
        """
        te = (self.headers.get("Transfer-Encoding") or "").lower()
        if "chunked" in te:
            out = bytearray()
            while True:
                line = self.rfile.readline()
                if not line:
                    break
                line = line.strip()
                if not line:
                    continue
                try:
                    size = int(line.split(b";")[0], 16)
                except ValueError:
                    break
                if size == 0:
                    # 消费 trailer 行直到空行
                    while True:
                        tail = self.rfile.readline()
                        if not tail or tail.strip() == b"":
                            break
                    break
                data = bytearray()
                remaining = size
                while remaining > 0:
                    chunk = self.rfile.read(remaining)
                    if not chunk:
                        return bytes(out)
                    data += chunk
                    remaining -= len(chunk)
                out += data
                self.rfile.read(2)  # 消费 chunk 后的 CRLF
            return bytes(out)
        n = int(self.headers.get("Content-Length", 0) or 0)
        return self.rfile.read(n) if n else b""

    def do_GET(self):
        p = self.path.split("?")[0].rstrip("/")
        if p in ("", "/health"):
            self._send(200, {"status": "ok", "proxy": "docproxy",
                             "mode": PROXY_MODE, "upstream": UPSTREAM})
        elif p in ("/files", "/v1/files"):
            self._send(200, load_index()["files"])
        elif p in ("/v1/models", "/models"):
            forward_to_upstream(self,"GET", "/v1/models")
        else:
            self._send(404, {"error": "not found", "path": self.path})

    def do_POST(self):
        raw = self._read_body()
        try:
            body = json.loads(raw) if raw else {}
        except json.JSONDecodeError:
            body = {}
        p = self.path.split("?")[0].rstrip("/")

        # ★ 2026-10-03 修复：剥离孤儿 media 占位符，避免 mmproj tokenizer 400
        #   （纯文本请求残留 <__media_xxx__> 却未附图 → 上游 tokenize 失败）。
        #   仅当整请求无图时才剥离；真多模态请求原样放行。
        if strip_stray_media_markers(body):
            raw = json.dumps(body).encode()

        if p in ("/v1/chat/completions", "/chat/completions"):
            # ★ 2026-10-03 头部裁剪：支持 X-Trim-Prompt 头动态剔除系统提示中的低频信息
            #   支持的值：
            #     connectors — 移除连接器状态（~2000 字符）
            #     automations — 移除自动化指令（~1500 字符）
            #     memory — 移除 workspace/用户/云端内存（~3000 字符）
            #     all — 全部移除（仅保留 SOUL.md + IDENTITY.md 核心身份）
            #   用法：curl -H "X-Trim-Prompt: all" ...
            trim_header = self.headers.get("X-Trim-Prompt", "").lower()
            if trim_header and isinstance(body, dict) and raw:
                msgs = body.get("messages", [])
                if msgs and msgs[0].get("role") == "system":
                    sys_content = msgs[0].get("content", "")
                    if isinstance(sys_content, str):
                        orig_len = len(sys_content)
                        if "connectors" in trim_header:
                            sys_content = re.sub(
                                r'(?:## 连接器状态|Connector Status|Connector Status\n)(?:[^#]|#(?! ))*(?=\n## |\Z)',
                                '', sys_content, flags=re.DOTALL
                            )
                        if "automations" in trim_header:
                            sys_content = re.sub(
                                r'(?:## 自动化|Automations|Automations\n)(?:[^#]|#(?! ))*(?=\n## |\Z)',
                                '', sys_content, flags=re.DOTALL
                            )
                        if "memory" in trim_header:
                            sys_content = re.sub(
                                r'(?:## 记忆|Workspace Memory|Cloud Memory|User-level Local|<memory>|</memory>)(?:[^#]|#(?! ))*(?=\n## |\Z)',
                                '', sys_content, flags=re.DOTALL
                            )
                        if "all" in trim_header:
                            # 保留 SOUL.md + IDENTITY.md，移除所有其他段落
                            parts = re.split(r'(?:^|\n)## (SOUL|IDENTITY)\.', sys_content)
                            if len(parts) > 2:
                                sys_content = parts[0] + '## SOUL.' + parts[1] + '## IDENTITY.' + parts[2]
                            else:
                                # 简化：只保留标题头
                                lines = sys_content.splitlines()
                                keep = []
                                for ln in lines:
                                    if 'SOUL' in ln or 'IDENTITY' in ln or ln.startswith('---'):
                                        keep.append(ln)
                                sys_content = '\n'.join(keep)
                        new_len = len(sys_content)
                        if new_len < orig_len:
                            body["messages"][0]["content"] = sys_content
                            raw = json.dumps(body).encode()

            # 注入默认关闭思考链（客户端显式指定时不覆盖）
            if INJECT_NO_THINK and isinstance(body, dict) and raw:
                ctk = body.get("chat_template_kwargs")
                if not isinstance(ctk, dict):
                    ctk = {}
                if "enable_thinking" not in ctk:
                    ctk["enable_thinking"] = False
                    body["chat_template_kwargs"] = ctk
                    raw = json.dumps(body).encode()

            # disclosure 模式 + 非流式 → 渐进式披露工具循环
            if PROXY_MODE == "disclosure" and not body.get("stream"):
                try:
                    self._send(200, handle_chat(body))
                except Exception as e:
                    self._send(500, {"error": str(e)})
            else:
                # passthrough 模式，或 disclosure 下的流式请求 → 原样透传（含 SSE）
                forward_to_upstream(self,"POST", "/v1/chat/completions", raw,
                                        stream=bool(body.get("stream")))
        elif p.startswith("/v1/") or p in ("/embeddings", "/completions"):
            forward_to_upstream(self,"POST", self.path.split("?")[0], raw)
        else:
            self._send(404, {"error": "not found", "path": self.path})


# ---------- CLI ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["serve", "add", "ls", "ask"])
    ap.add_argument("args", nargs="*")
    ap.add_argument("--port", type=int, default=8081)
    ap.add_argument("--upstream", default=None)
    a = ap.parse_args()
    global UPSTREAM
    if a.upstream:
        UPSTREAM = a.upstream

    if a.cmd == "add":
        for p in a.args:
            register(p)
    elif a.cmd == "ls":
        files = load_index()["files"]
        for k, v in files.items():
            print(f"id={k}  {v['name']}  {v['size_kb']} KB  {v.get('chars',0)} 字符")
        if not files:
            print("(空) 用 add 注册文件")
    elif a.cmd == "ask":
        q = " ".join(a.args) or "文档第一个标题是什么？"
        t0 = time.time()
        r = handle_chat({"messages": [{"role": "user", "content": q}],
                         "max_tokens": 512, "temperature": 0.2})
        s = r["_stats"]
        print(f"--- 答案 ---\n{r['choices'][0]['message']['content']}")
        print(f"--- 统计: 轮次={s['rounds']} 工具调用={len(s['tool_calls'])} "
              f"累计 prompt_tokens={s['prompt_tokens']} 耗时={time.time()-t0:.1f}s")
        for t in s["tool_calls"]:
            print(f"  {t}")
    else:
        print(f"docproxy :{a.port} -> {UPSTREAM}  mode={PROXY_MODE}")
        print(build_catalog()[:400] or "(索引为空)")
        ThreadingHTTPServer(("127.0.0.1", a.port), Handler).serve_forever()


if __name__ == "__main__":
    main()
